import io
import re
from collections import Counter

import pandas as pd
import plotly.express as px
import spacy
import streamlit as st
from docx import Document
from pypdf import PdfReader

st.set_page_config(
    page_title="LexiTag | POS Intelligence Studio",
    page_icon="🧠",
    layout="wide",
    initial_sidebar_state="expanded",
)

# ---------- Styling ----------
st.markdown("""
<style>
@import url('https://fonts.googleapis.com/css2?family=DM+Sans:wght@400;500;600;700&family=Space+Grotesk:wght@500;600;700&display=swap');
html, body, [class*="css"] { font-family: 'DM Sans', sans-serif; }
.stApp { background: radial-gradient(circle at 85% 2%, rgba(87, 75, 255, .12), transparent 27%), #0b1020; color: #edf0ff; }
[data-testid="stSidebar"] { background: #10172b; border-right: 1px solid #28314b; }
.hero { padding: 26px 30px; border: 1px solid #2a3553; border-radius: 22px;
        background: linear-gradient(120deg, rgba(91,75,255,.22), rgba(25,190,180,.08)), #111a30; margin-bottom: 22px; }
.eyebrow { color: #8ce6dd; font-size: .76rem; letter-spacing: .17em; font-weight: 700; text-transform: uppercase; }
.hero h1 { font-family: 'Space Grotesk', sans-serif; font-size: 2.55rem; line-height: 1.1; margin: 10px 0; color: #f5f6ff; }
.hero p { color: #b7c2df; max-width: 820px; margin-bottom: 0; font-size: 1.02rem; }
.section-title { font-family: 'Space Grotesk', sans-serif; font-size: 1.25rem; font-weight: 700; margin: 12px 0 8px; }
.metric-card { background: #131d34; border: 1px solid #2a3553; border-radius: 16px; padding: 18px 20px; }
.metric-label { color: #9eaccb; font-size: .82rem; }
.metric-value { font-family: 'Space Grotesk', sans-serif; color: #f6f7ff; font-size: 1.8rem; font-weight: 700; margin-top: 4px; }
.mini-note { color: #aab7d6; font-size: .88rem; }
.tag-chip { display: inline-block; border: 1px solid #38496c; background: #172541; color: #b8fff2; padding: 5px 9px; margin: 3px; border-radius: 8px; font-size: .82rem; }
div[data-testid="stDataFrame"] { border: 1px solid #283653; border-radius: 12px; overflow: hidden; }
.stButton > button, .stDownloadButton > button { border-radius: 10px; font-weight: 700; }
hr { border-color: #283653; }
footer { visibility: hidden; }
</style>
""", unsafe_allow_html=True)

@st.cache_resource(show_spinner="Loading English NLP pipeline...")
def load_nlp():
    return spacy.load("en_core_web_sm")

def extract_text(uploaded_file):
    name = uploaded_file.name.lower()
    raw = uploaded_file.getvalue()
    if name.endswith(".txt") or name.endswith(".md"):
        return raw.decode("utf-8", errors="replace")
    if name.endswith(".docx"):
        doc = Document(io.BytesIO(raw))
        parts = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text.strip() for cell in row.cells))
        return "\n".join(parts)
    if name.endswith(".pdf"):
        reader = PdfReader(io.BytesIO(raw))
        return "\n".join(page.extract_text() or "" for page in reader.pages)
    raise ValueError("Unsupported file type. Upload PDF, DOCX, TXT, or Markdown.")

def make_annotated_text(doc):
    lines = []
    for token in doc:
        if token.is_space:
            continue
        lines.append(f"{token.text}  [{token.pos_} | {token.tag_} | {spacy.explain(token.tag_) or token.pos_}]")
    return "\n".join(lines)

def make_tag_table(doc):
    rows = []
    for i, token in enumerate(doc):
        if token.is_space:
            continue
        rows.append({
            "Index": i,
            "Token": token.text,
            "Lemma": token.lemma_,
            "POS": token.pos_,
            "Fine Tag": token.tag_,
            "Tag Meaning": spacy.explain(token.tag_) or token.pos_,
            "Dependency": token.dep_,
            "Head": token.head.text,
            "Is Stopword": "Yes" if token.is_stop else "No",
            "Is Punctuation": "Yes" if token.is_punct else "No",
        })
    return pd.DataFrame(rows)

def metric_card(label, value):
    st.markdown(f'<div class="metric-card"><div class="metric-label">{label}</div><div class="metric-value">{value}</div></div>', unsafe_allow_html=True)

# ---------- Sidebar ----------
with st.sidebar:
    st.markdown("## 🧠 LexiTag")
    st.caption("POS Intelligence Studio")
    st.markdown("---")
    st.markdown("### Workspace")
    source_mode = st.radio("Input source", ["Upload document", "Paste text"], index=0)
    st.markdown("---")
    st.markdown("### Analysis settings")
    show_punct = st.toggle("Include punctuation in results", value=False)
    show_stopwords = st.toggle("Include stopwords in results", value=True)
    st.markdown("---")
    st.markdown("**Pipeline**")
    st.markdown("`Document → Text extraction → spaCy → POS tags → Insights`")
    st.caption("Designed for educational NLP analysis. Do not upload confidential documents to a hosted deployment.")

# ---------- Hero ----------
st.markdown("""
<div class="hero">
  <div class="eyebrow">NATURAL LANGUAGE PROCESSING · DOCUMENT ANALYTICS</div>
  <h1>LexiTag <span style="color:#8ce6dd">Studio</span></h1>
  <p>Turn documents into structured linguistic insights. Upload a file, inspect token-level Part-of-Speech tags, explore grammar patterns, and export your analysis in seconds.</p>
</div>
""", unsafe_allow_html=True)

st.markdown("#### Start an analysis")
uploaded_file = None
pasted_text = ""
if source_mode == "Upload document":
    uploaded_file = st.file_uploader("Choose a document", type=["pdf", "docx", "txt", "md"], help="Supported: PDF, Word (.docx), plain text (.txt), Markdown (.md)")
    st.caption("PDFs with selectable text work best. Scanned PDFs require OCR, which is not included in this version.")
else:
    pasted_text = st.text_area("Paste text to analyse", height=190, placeholder="Paste a paragraph, article, or sample passage here...")

if st.button("✨ Analyse document", type="primary", use_container_width=True):
    try:
        text = extract_text(uploaded_file) if uploaded_file else pasted_text
        if not text or not text.strip():
            st.session_state.pop("analysis_text", None)
            st.warning("Please upload a document or paste some text first.")
        elif len(text) > 1_000_000:
            st.warning("This document is larger than 1 MB of extracted text. Please use a smaller document for this demo.")
        else:
            nlp = load_nlp()
            doc = nlp(text)
            st.session_state["analysis_text"] = text
            st.session_state["analysis_doc_bytes"] = uploaded_file.name if uploaded_file else "Pasted text"
            st.session_state["analysis_rows"] = make_tag_table(doc).to_dict("records")
            st.session_state["analysis_annotated"] = make_annotated_text(doc)
            st.session_state["analysis_text_preview"] = text[:5000]
            st.session_state["analysis_summary"] = {
                "token_count": sum(1 for t in doc if not t.is_space),
                "word_count": sum(1 for t in doc if t.is_alpha),
                "sentence_count": sum(1 for _ in doc.sents),
                "unique_lemmas": len({t.lemma_.lower() for t in doc if t.is_alpha}),
            }
            st.success("Analysis complete. Explore the results below.")
    except Exception as e:
        st.error(f"Could not analyse this input: {e}")
        st.info("If this is a fresh setup, install dependencies from requirements.txt and make sure the spaCy English model is available.")

if "analysis_rows" not in st.session_state:
    st.markdown("---")
    st.markdown("### What you can explore")
    c1, c2, c3 = st.columns(3)
    with c1:
        st.markdown("**01 · Token explorer**")
        st.write("Inspect each token, its lemma, universal POS, fine-grained tag, and grammatical dependency.")
    with c2:
        st.markdown("**02 · Linguistic dashboard**")
        st.write("Review word and sentence counts, vocabulary variety, and the most frequent POS categories.")
    with c3:
        st.markdown("**03 · Export-ready results**")
        st.write("Download token-level annotations as CSV or the full tagged text as a TXT file.")
    st.stop()

# Recreate spaCy doc from cached text for charts and sentences.
nlp = load_nlp()
doc = nlp(st.session_state["analysis_text"])
full_df = pd.DataFrame(st.session_state["analysis_rows"])
view_df = full_df.copy()
if not show_punct:
    view_df = view_df[view_df["Is Punctuation"] == "No"]
if not show_stopwords:
    view_df = view_df[view_df["Is Stopword"] == "No"]

summary = st.session_state["analysis_summary"]
pos_counts = full_df[full_df["Is Punctuation"] == "No"]["POS"].value_counts().reset_index()
pos_counts.columns = ["POS", "Count"]
tag_counts = full_df[full_df["Is Punctuation"] == "No"]["Fine Tag"].value_counts().head(12).reset_index()
tag_counts.columns = ["Fine Tag", "Count"]
word_tokens = [t for t in doc if t.is_alpha]
word_freq = Counter(t.lemma_.lower() for t in word_tokens if not t.is_stop and len(t.lemma_) > 2)
common_words = pd.DataFrame(word_freq.most_common(12), columns=["Lemma", "Frequency"])

st.markdown("---")
st.markdown("### Analysis overview")
st.caption(f"Source: {st.session_state.get('analysis_doc_bytes', 'Document')} · English pipeline · spaCy POS tagger")
m1, m2, m3, m4 = st.columns(4)
with m1: metric_card("TOKENS", f"{summary['token_count']:,}")
with m2: metric_card("WORDS", f"{summary['word_count']:,}")
with m3: metric_card("SENTENCES", f"{summary['sentence_count']:,}")
with m4: metric_card("UNIQUE LEMMAS", f"{summary['unique_lemmas']:,}")

tab1, tab2, tab3, tab4 = st.tabs(["📊 Insights", "🔎 Token explorer", "📝 Document preview", "⬇️ Export"])

with tab1:
    left, right = st.columns([1.15, 1])
    with left:
        st.markdown('<div class="section-title">Part-of-Speech distribution</div>', unsafe_allow_html=True)
        fig = px.bar(pos_counts.sort_values("Count", ascending=True), x="Count", y="POS", orientation="h",
                     text="Count", color="Count", color_continuous_scale=["#2dd4bf", "#7468ff"])
        fig.update_layout(template="plotly_dark", paper_bgcolor="rgba(0,0,0,0)", plot_bgcolor="rgba(0,0,0,0)",
                          margin=dict(l=10,r=10,t=10,b=10), coloraxis_showscale=False, xaxis_title="Token count", yaxis_title="")
        st.plotly_chart(fig, use_container_width=True)
    with right:
        st.markdown('<div class="section-title">Most frequent fine-grained tags</div>', unsafe_allow_html=True)
        fig2 = px.pie(tag_counts, names="Fine Tag", values="Count", hole=.58,
                      color_discrete_sequence=px.colors.qualitative.Bold)
        fig2.update_layout(template="plotly_dark", paper_bgcolor="rgba(0,0,0,0)", margin=dict(l=5,r=5,t=10,b=10),
                           legend_title_text="Fine tag")
        st.plotly_chart(fig2, use_container_width=True)
    st.markdown('<div class="section-title">Most frequent content lemmas</div>', unsafe_allow_html=True)
    if len(common_words):
        fig3 = px.bar(common_words.sort_values("Frequency"), x="Frequency", y="Lemma", orientation="h",
                      text="Frequency", color="Frequency", color_continuous_scale=["#1d9e91", "#7468ff"])
        fig3.update_layout(template="plotly_dark", paper_bgcolor="rgba(0,0,0,0)", plot_bgcolor="rgba(0,0,0,0)",
                           margin=dict(l=10,r=10,t=10,b=10), coloraxis_showscale=False, xaxis_title="Frequency", yaxis_title="")
        st.plotly_chart(fig3, use_container_width=True)
    st.markdown("**Quick POS guide**")
    st.markdown('<span class="tag-chip">NOUN · person/place/thing</span><span class="tag-chip">VERB · action/state</span><span class="tag-chip">ADJ · describes a noun</span><span class="tag-chip">ADV · modifies verb/adjective</span><span class="tag-chip">PRON · replaces a noun</span><span class="tag-chip">DET · the/a/this</span>', unsafe_allow_html=True)

with tab2:
    st.markdown('<div class="section-title">Token-level annotations</div>', unsafe_allow_html=True)
    c1, c2 = st.columns([1.2, 1])
    with c1:
        search_term = st.text_input("Search token or lemma", placeholder="e.g. running")
    with c2:
        pos_options = ["All"] + sorted(full_df["POS"].dropna().unique().tolist())
        chosen_pos = st.selectbox("Filter by POS", pos_options)
    if search_term.strip():
        mask = view_df["Token"].str.contains(search_term, case=False, na=False) | view_df["Lemma"].str.contains(search_term, case=False, na=False)
        view_df = view_df[mask]
    if chosen_pos != "All":
        view_df = view_df[view_df["POS"] == chosen_pos]
    st.caption(f"Showing {len(view_df):,} of {len(full_df):,} token records. Toggle punctuation and stopwords in the sidebar.")
    st.dataframe(view_df, use_container_width=True, hide_index=True, height=430)
    st.markdown("**Tag legend**")
    st.write("`POS` = universal category; `Fine Tag` = more specific English grammar tag; `Dependency` = grammatical relationship to another token; `Lemma` = dictionary/base form.")

with tab3:
    st.markdown('<div class="section-title">Extracted text</div>', unsafe_allow_html=True)
    st.text_area("Document text (preview, up to 5,000 characters)", value=st.session_state["analysis_text_preview"], height=220, disabled=True)
    with st.expander("View sentence segmentation"):
        for i, sent in enumerate(doc.sents, start=1):
            st.markdown(f"**Sentence {i}**")
            st.write(sent.text)
    with st.expander("View POS-tagged text"):
        st.text(st.session_state["analysis_annotated"][:30000])
        if len(st.session_state["analysis_annotated"]) > 30000:
            st.caption("Preview limited to 30,000 characters; download the full annotation in Export.")

with tab4:
    st.markdown('<div class="section-title">Download your analysis</div>', unsafe_allow_html=True)
    st.write("Export the complete token table for spreadsheet work, or a readable text file containing each token and its tags.")
    csv_bytes = full_df.to_csv(index=False).encode("utf-8")
    annotated_bytes = st.session_state["analysis_annotated"].encode("utf-8")
    original_bytes = st.session_state["analysis_text"].encode("utf-8")
    c1, c2, c3 = st.columns(3)
    with c1:
        st.download_button("⬇️ Download POS table (CSV)", data=csv_bytes, file_name="pos_tagging_results.csv", mime="text/csv", use_container_width=True)
    with c2:
        st.download_button("⬇️ Download tagged text (TXT)", data=annotated_bytes, file_name="pos_tagged_text.txt", mime="text/plain", use_container_width=True)
    with c3:
        st.download_button("⬇️ Download extracted text (TXT)", data=original_bytes, file_name="extracted_document.txt", mime="text/plain", use_container_width=True)
    st.markdown("---")
    st.markdown("**Suggested project demonstration**")
    st.markdown("1. Upload a sample `.docx`, `.pdf`, or `.txt` file.  \n2. Click **Analyse document**.  \n3. Show POS charts and filter tokens by category.  \n4. Download the CSV and explain how POS tagging supports downstream NLP tasks.")

st.markdown("---")
st.caption("LexiTag Studio · NLP Mini Project · Built with Streamlit, spaCy, pandas and Plotly")
